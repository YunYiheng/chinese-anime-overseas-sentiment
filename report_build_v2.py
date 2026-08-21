# -*- coding: utf-8 -*-
"""
report_build_v2.py —— 将 report_draft_v2.md 排版为最终 Word。

  1. 解析 markdown（标题/段落/表格/引用块/列表/图片引用），转为 python-docx 文档。
  2. 图表编号连续化：图 1–11、表 1–9（§2.1/§2.2 两张表补为表 1/表 2）。
  3. 中文字体（宋体/黑体）+ 英文数字字体（Times New Roman）分离设置。
  4. 目录字段（TOC \o "1-2" \h，带超链接，需在 Word 中 F9 更新）。
  5. 表格跨页处理：表头重复 + 单元格禁止跨行拆分。
  6. 图片居中插入 + 题注。

运行：python report_build_v2.py
输出：国产动画海外舆情分析报告.docx
"""
import re
import sys

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

MD = "report_draft_v2.md"
OUT = "国产动画海外舆情分析报告.docx"

# --------------------------------------------------------------------- 图表元数据
FIGURES = [
    ("cover_matrix.png", "作品 × 平台评论覆盖矩阵（格内为样本数）"),
    ("series_platform_dist.png", "Series 各平台主样本量（排除 other）"),
    ("series_time_dist.png", "Series 评论发布时间分布（按年月）"),
    ("series_sentiment_dist.png", "Series 三口径情感分布（%）"),
    ("series_topic_overall_distribution.png", "Series 主题命中分布"),
    ("07_residual_series.png", "Series 平台 × 情感标准化残差热力图"),
    ("film_platform_dist.png", "Film 各平台主样本量（排除 other）"),
    ("film_time_dist.png", "Film 评论发布时间分布（按年月）"),
    ("film_sentiment_dist.png", "Film 三口径情感分布（%）"),
    ("film_topic_overall_distribution.png", "Film 主题命中分布"),
    ("07_residual_film.png", "Film 平台 × 情感标准化残差热力图"),
]
FIG_PATH = {
    "cover_matrix.png": "cover_matrix.png",
    "series_platform_dist.png": "series_platform_dist.png",
    "film_platform_dist.png": "film_platform_dist.png",
    "series_sentiment_dist.png": "series_sentiment_dist.png",
    "film_sentiment_dist.png": "film_sentiment_dist.png",
    "series_topic_overall_distribution.png": "storage2/series_pipeline/series_topic_overall_distribution.png",
    "film_topic_overall_distribution.png": "storage2/film_pipeline/film_topic_overall_distribution.png",
    "series_time_dist.png": "series_time_dist.png",
    "film_time_dist.png": "film_time_dist.png",
    "07_residual_series.png": "07_residual_series.png",
    "07_residual_film.png": "07_residual_film.png",
}
FIG_WIDTH = {
    "cover_matrix.png": Cm(16.5),
    "07_residual_series.png": Cm(12.0),
    "07_residual_film.png": Cm(12.0),
}
FIG_NUM = {name: i + 1 for i, (name, _) in enumerate(FIGURES)}
FIG_TITLE = dict(FIGURES)

TABLE_CAPTIONS = [
    "Series 各作品主样本量",
    "Film 各作品主样本量",
    "Series 三口径情感分布（%）",
    "Series 各平台正面率与内容形态",
    "Film 三口径情感分布（%）",
    "Film 各平台正面率与内容形态",
    "Series vs Film 情感交叉表（YouTube + X）",
    "同一部《哪吒》各平台正面率",
    "平台评论长度与主题命中率",
]

CN_BODY = "宋体"
CN_HEAD = "黑体"
EN_BODY = "Times New Roman"
GRAY = RGBColor(0x55, 0x55, 0x55)
DARK = RGBColor(0x33, 0x33, 0x33)

# --------------------------------------------------------------------- 工具函数
def set_run_font(run, cn=CN_BODY, en=EN_BODY, size=None, bold=None, color=None):
    run.font.name = en
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), cn)


def set_style_font(style, cn, en, size, bold=None, color=None):
    style.font.name = en
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), cn)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+?`)")

def add_rich(par, text, size=None, color=None, bold=False):
    """解析 **粗体** 与 `行内代码`，写入段落。"""
    for tok in INLINE_RE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = par.add_run(tok[2:-2])
            set_run_font(r, size=size, bold=True, color=color)
        elif tok.startswith("`") and tok.endswith("`"):
            r = par.add_run(tok[1:-1])
            set_run_font(r, en="Consolas", size=(size or 10), color=DARK)
        else:
            r = par.add_run(tok)
            set_run_font(r, size=size, bold=bold, color=color)


def add_toc(doc):
    h = doc.add_heading(level=1)
    r = h.add_run("目录")
    set_run_font(r, cn=CN_HEAD, size=15, bold=True, color=RGBColor(0, 0, 0))
    h.paragraph_format.space_after = Pt(6)

    par = doc.add_paragraph()
    run = par.add_run()
    fld = OxmlElement("w:fldChar"); fld.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-2" \h \z \u'
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = "（请在 Word 中右键选择“更新域”以生成目录）"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.append(fld); run._r.append(instr); run._r.append(sep); run._r.append(t); run._r.append(end)


def add_pagebreak(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run()
    br = OxmlElement("w:br"); br.set(qn("w:type"), "page")
    run._r.append(br)


def add_figure(doc, num, name):
    path = FIG_PATH[name]
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    width = FIG_WIDTH.get(name, Cm(14.0))
    try:
        run.add_picture(path, width=width)
    except FileNotFoundError:
        run.text = f"[缺图：{path}]"
        set_run_font(run, color=RGBColor(0xC0, 0x00, 0x00), size=9)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    r1 = cap.add_run(f"图 {num}　")
    set_run_font(r1, cn=CN_HEAD, size=9, bold=True, color=DARK)
    r2 = cap.add_run(FIG_TITLE[name])
    set_run_font(r2, size=9, color=DARK)


def _set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true")
    trPr.append(th)


def _set_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    cs = OxmlElement("w:cantSplit")
    trPr.append(cs)


def add_table(doc, header, rows, caption):
    cap = doc.add_paragraph()
    cap.paragraph_format.space_before = Pt(6)
    cap.paragraph_format.space_after = Pt(2)
    cap.paragraph_format.keep_with_next = True
    r1 = cap.add_run(caption)
    set_run_font(r1, cn=CN_HEAD, size=9, bold=True, color=DARK)

    tbl = doc.add_table(rows=1 + len(rows), cols=len(header))
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = True

    hdr = tbl.rows[0]
    _set_repeat_header(hdr)
    _set_cant_split(hdr)
    for j, htxt in enumerate(header):
        cell = hdr.cells[j]
        cell.paragraphs[0].text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(htxt.replace("**", ""))
        set_run_font(r, cn=CN_HEAD, size=9, bold=True)

    for i, row in enumerate(rows):
        tr = tbl.rows[i + 1]
        _set_cant_split(tr)
        for j, val in enumerate(row):
            cell = tr.cells[j]
            cell.paragraphs[0].text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_rich(p, val, size=9)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(2)
    sp.paragraph_format.space_before = Pt(0)
    r = sp.add_run("")
    set_run_font(r, size=2)


def extract_figures(text):
    """提取段落中的图片引用，返回 (清洗后文本, [图编号...])。"""
    figs = [FIG_NUM[m.group(1)] for m in re.finditer(r"`([^`]+\.png)`", text)]
    text = re.sub(r"（图\s*\d+\s*[，,]\s*`[^`]+\.png`）", "", text)
    text = re.sub(r"（`[^`]+\.png`）", "", text)
    text = re.sub(r"（`[^`]+\.png`[，,]", "（", text)
    return text.strip(), figs


def add_blockquote_line(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.space_after = Pt(2)
    m = re.match(r"^(EN|中|理由|记录编号)[：:]\s*", text)
    if m:
        r0 = p.add_run(m.group(0))
        set_run_font(r0, cn=CN_HEAD, size=10, bold=True, color=GRAY)
        rest = text[m.end():]
    else:
        rest = text
    add_rich(p, rest, size=10, color=GRAY)


def is_table_separator(s):
    return bool(re.match(r"^\s*\|[\s:|-]+\|?\s*$", s))


def parse_table(block_lines):
    header = [c.strip() for c in block_lines[0].strip().strip("|").split("|")]
    rows = []
    for ln in block_lines[1:]:
        if is_table_separator(ln):
            continue
        rows.append([c.strip() for c in ln.strip().strip("|").split("|")])
    return header, rows


# --------------------------------------------------------------------- 主流程
def main():
    doc = Document()

    for sec in doc.sections:
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        sec.top_margin = Cm(2.0)
        sec.bottom_margin = Cm(2.0)
        sec.left_margin = Cm(2.2)
        sec.right_margin = Cm(2.2)

    set_style_font(doc.styles["Normal"], CN_BODY, EN_BODY, 11)
    set_style_font(doc.styles["Heading 1"], CN_HEAD, EN_BODY, 16, bold=True, color=RGBColor(0, 0, 0))
    set_style_font(doc.styles["Heading 2"], CN_HEAD, EN_BODY, 13, bold=True, color=RGBColor(0, 0, 0))

    with open(MD, encoding="utf-8") as f:
        lines = f.read().split("\n")

    table_counter = 0
    toc_added = False
    i = 0
    n = len(lines)

    while i < n:
        line = lines[i].rstrip()
        s = line.strip()

        if not s:
            i += 1
            continue

        if s == "---":
            i += 1
            continue

        # 章节标题（## 数字/附录）
        if s.startswith("## "):
            title = s[3:].strip()
            if title.startswith("——"):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(10)
                add_rich(p, title, size=13)
                for r in p.runs:
                    set_run_font(r, cn=CN_HEAD, size=13)
            else:
                if not toc_added:
                    # 首个章节前：插入目录
                    add_pagebreak(doc)
                    add_toc(doc)
                    add_pagebreak(doc)
                    toc_added = True
                else:
                    add_pagebreak(doc)
                h = doc.add_heading(level=1)
                h.paragraph_format.space_before = Pt(6)
                h.paragraph_format.space_after = Pt(8)
                add_rich(h, title, size=16)
                for r in h.runs:
                    set_run_font(r, cn=CN_HEAD, en=EN_BODY, size=16, bold=True, color=RGBColor(0, 0, 0))
            i += 1
            continue

        # 文档主标题（单个 #）
        if s.startswith("# "):
            title = s[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(60)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(title)
            set_run_font(r, cn=CN_HEAD, size=22, bold=True, color=RGBColor(0, 0, 0))
            i += 1
            continue

        # 小节标题（### ）
        if s.startswith("### "):
            title = s[4:].strip()
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(6)
            add_rich(h, title, size=13)
            for r in h.runs:
                set_run_font(r, cn=CN_HEAD, en=EN_BODY, size=13, bold=True, color=RGBColor(0, 0, 0))
            i += 1
            continue

        # 表格
        if s.startswith("|"):
            block = []
            while i < n and lines[i].strip().startswith("|"):
                block.append(lines[i].rstrip())
                i += 1
            header, rows = parse_table(block)
            caption = f"表 {table_counter + 1}　{TABLE_CAPTIONS[table_counter]}"
            table_counter += 1
            add_table(doc, header, rows, caption)
            continue

        # 引用块
        if s.startswith(">"):
            txt = s[1:].strip()
            if txt.startswith("本稿为正文草稿"):
                i += 1
                continue
            if txt.startswith("数据来源"):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(2)
                add_rich(p, txt, size=10, color=GRAY)
            else:
                add_blockquote_line(doc, txt)
            i += 1
            continue

        # 无序列表
        if s.startswith("- "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run("• ")
            set_run_font(r, size=11)
            add_rich(p, s[2:].strip(), size=11)
            i += 1
            continue

        # 有序列表
        m_ol = re.match(r"^(\d+)\.\s+(.*)$", s)
        if m_ol:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(m_ol.group(1) + ". ")
            set_run_font(r, size=11, bold=True)
            add_rich(p, m_ol.group(2), size=11)
            i += 1
            continue

        # 段落（含图片引用）
        if "草稿完" in s:
            i += 1
            continue

        text, figs = extract_figures(s)
        if text:
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.3
            p.paragraph_format.space_after = Pt(4)
            add_rich(p, text, size=11)
        for fnum in figs:
            add_figure(doc, fnum, FIGURES[fnum - 1][0])
        i += 1

    doc.save(OUT)
    print(f"[OK] 已生成 {OUT}")
    print(f"     表格 {table_counter} 张（应为 9）")


if __name__ == "__main__":
    main()
