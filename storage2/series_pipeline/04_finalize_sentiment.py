# -*- coding: utf-8 -*-
"""
04_finalize_sentiment.py
========================
情感人工复核结果回填 + 最终情感汇总（三统计口径版）。

数据流：
    03_sentiment_all.csv             全量基线（含模型 sentiment_pred，manual 列空白占位）
        +
    复核表 xlsx（核心复核清单 + 补充抽审清单，人工已填 manual_sentiment）
        ↓
    04 最终 csv（全量 + 回填后的 sentiment_final + 新增列）
    汇总 xlsx（含三口径对比 / 平台准入 / 一致率等）

回填口径：
    manual_sentiment 非空 → sentiment_final = manual_sentiment
    否则             → sentiment_final = sentiment_pred（模型结果）

【三统计口径】（本版核心改动，取代旧的"自然/平台均衡"两口径）：
    1. pooled              ：全部有效评论（排除 other）直接汇总，不看平台。
    2. platform_balanced   ：满足准入门槛的平台做"平台内占比→等权平均"。
                             ★ 这是正文主口径。
    3. all_platform_equal  ：所有平台等权平均（含小平台），仅作敏感性分析。

【平台准入规则】（自动判断，不硬编码平台列表）：
    一个平台要进入 platform_balanced，须同时满足：
      - 样本量 n >= 30（有效评论数，排除 other）
      - 覆盖作品数 >= 3（该平台里出现过的不同作品数）
    小平台（如 AniList n=27、或只覆盖 1 部作品的 Reddit@film）不进入正文均衡，
    但保留在"数据 + 平台级 EDA + all_platform_equal 敏感性"里。

用法：
    python 04_finalize_sentiment.py --label Series \
        --csv 03_sentiment_all.csv --xlsx series_sentiment_review_core.xlsx \
        --out-csv series_sentiment_final.csv --out-xlsx series_sentiment_final_summary.xlsx

    python 04_finalize_sentiment.py --label Film \
        --csv 03_sentiment_all.csv --xlsx film_sentiment_review_core.xlsx \
        --out-csv film_sentiment_final.csv --out-xlsx film_sentiment_final_summary.xlsx
"""

import argparse
import sys

import numpy as np
import pandas as pd

sys.stdout.reconfigure(encoding="utf-8")

# ---------------- 常量 ----------------
# manual_sentiment 允许值
ALLOWED_MANUAL = {"positive", "neutral", "negative", "mixed", "other"}
# 主情感档（不含 other）：比例统计的主分母只用这四档
EMOTIONS = ["positive", "neutral", "mixed", "negative"]

# 平台准入门槛
MIN_PLATFORM_N = 30       # 样本量门槛（有效评论数）
MIN_PLATFORM_WORKS = 3    # 覆盖作品数门槛

# 明细表内容预览截断长度
CONTENT_PREVIEW = 120


# ============================================================
# 一、读取 + 合并两张复核表
# ============================================================
def load_review_tables(xlsx_path):
    """读取复核 xlsx 的『核心复核清单』『补充抽审清单』并合并，打队列标记。"""
    xl = pd.ExcelFile(xlsx_path)
    core = xl.parse("核心复核清单").copy()
    extra = xl.parse("补充抽审清单").copy()
    core["review_queue_group"] = "核心复核清单"
    extra["review_queue_group"] = "补充抽审清单"
    review = pd.concat([core, extra], ignore_index=True)
    print(f"[1/7] 读取复核表: 核心 {len(core)} + 补充 {len(extra)} = {len(review)} 条")
    return review


# ============================================================
# 二、校验
# ============================================================
def validate_review(review, main_df, lenient):
    """四类校验：record_id 不重复 / 能匹配 / manual 值域 / manual 不空白。"""
    errors, warnings = [], []

    dup = review.loc[review["record_id"].duplicated(keep=False), "record_id"]
    if len(dup):
        errors.append(f"record_id 重复 {dup.unique()[:10]}")
    else:
        print("[2/7] record_id 无重复 ✓")

    known = set(main_df["record_id"])
    miss = review.loc[~review["record_id"].isin(known), "record_id"]
    if len(miss):
        errors.append(f"{len(miss)} 个 record_id 在 03 中不存在: {miss.unique()[:10]}")
    else:
        print("[2/7] record_id 全部能在 03 中匹配 ✓")

    manual_norm = review["manual_sentiment"].fillna("").astype(str).str.strip()
    bad_val = review.loc[~manual_norm.isin(ALLOWED_MANUAL) & manual_norm.ne(""), "record_id"]
    if len(bad_val):
        errors.append(f"{len(bad_val)} 条 manual_sentiment 非法值 "
                      f"{review.loc[~manual_norm.isin(ALLOWED_MANUAL) & manual_norm.ne(''), 'manual_sentiment'].unique()[:10]}")
    else:
        print("[2/7] manual_sentiment 值域合法 ✓")

    blank = int(manual_norm.eq("").sum())
    if blank > 0:
        msg = f"{blank} 条 manual_sentiment 未填写，正式模式禁止用模型结果兜底。"
        if lenient:
            warnings.append(msg + "（lenient 降级用模型结果）")
        else:
            errors.append(msg)
    else:
        print(f"[2/7] {len(review)} 条 manual_sentiment 全部填写 ✓")

    for w in warnings:
        print(f"      ⚠ {w}")
    if errors:
        print("\n======== 校验失败，终止 ========")
        for e in errors:
            print(f"  ✗ {e}")
        sys.exit(1)
    return blank


# ============================================================
# 三、回填 + 最终化
# ============================================================
def merge_back(main_df, review):
    """把复核表的 manual 结果按 record_id 回填到全量。"""
    back = review[["record_id", "manual_sentiment", "manual_note",
                   "review_queue_group", "rule_hit"]].rename(
        columns={"rule_hit": "review_rule_hit"})
    out = main_df.drop(columns=["manual_sentiment", "manual_note", "sentiment_final"],
                       errors="ignore").merge(back, on="record_id", how="left")
    for c in ["manual_sentiment", "manual_note", "review_queue_group", "review_rule_hit"]:
        out[c] = out[c].fillna("")
    return out


def finalize(df):
    """生成 sentiment_final / manual_reviewed / label_changed。"""
    mask = df["manual_sentiment"] != ""
    df["sentiment_final"] = df["sentiment_pred"]
    df.loc[mask, "sentiment_final"] = df.loc[mask, "manual_sentiment"]
    df["manual_reviewed"] = np.where(mask, "是", "否")
    changed = np.where(df["manual_sentiment"] != df["sentiment_pred"], "是", "否")
    df["label_changed"] = pd.Series(changed, index=df.index).where(mask, "")
    return df


# ============================================================
# 四、平台准入 + 三口径（本版核心）
# ============================================================
def compute_platform_eligibility(main_df):
    """按 n>=30 且 覆盖作品>=3 自动判断每个平台是否进入正文均衡。"""
    rows = []
    for src, g in main_df.groupby("source"):
        n = len(g)
        works = g["anime_name"].nunique()
        rows.append({
            "平台": src,
            "样本量n": n,
            "覆盖作品数": works,
            "n>=30": n >= MIN_PLATFORM_N,
            "作品>=3": works >= MIN_PLATFORM_WORKS,
            "纳入均衡": (n >= MIN_PLATFORM_N) and (works >= MIN_PLATFORM_WORKS),
        })
    elig = pd.DataFrame(rows).sort_values("样本量n", ascending=False).reset_index(drop=True)
    return elig


def _platform_proportions(main_df, platforms):
    """对给定平台列表，算每个平台内四情感占比，返回 (platforms × EMOTIONS) 的 DataFrame。"""
    pivot = main_df[main_df["source"].isin(platforms)].pivot_table(
        index="source", columns="sentiment_final", values="record_id",
        aggfunc="count", fill_value=0)
    pivot = pivot.reindex(columns=EMOTIONS, fill_value=0)
    ratio = pivot.div(pivot.sum(axis=1), axis=0).fillna(0)
    return ratio


def compute_three_views(main_df):
    """计算三口径对比，返回 (对比表, 平台准入表, 准入平台集合)。"""
    n_main = len(main_df)
    elig = compute_platform_eligibility(main_df)
    eligible = set(elig.loc[elig["纳入均衡"], "平台"])
    all_platforms = list(main_df["source"].unique())

    # 1) pooled：全部有效评论直接汇总
    vc = main_df["sentiment_final"].value_counts().reindex(EMOTIONS, fill_value=0)
    pooled = vc / vc.sum() * 100

    # 2) platform_balanced：准入平台等权
    ratio_bal = _platform_proportions(main_df, eligible)
    balanced = ratio_bal.mean(axis=0) * 100 if len(ratio_bal) else pd.Series(np.nan, index=EMOTIONS)

    # 3) all_platform_equal：所有平台等权
    ratio_all = _platform_proportions(main_df, all_platforms)
    all_equal = ratio_all.mean(axis=0) * 100 if len(ratio_all) else pd.Series(np.nan, index=EMOTIONS)

    comp = pd.DataFrame({
        "pooled_全部汇总%": pooled.round(2),
        f"platform_balanced_{len(eligible)}平台等权%": balanced.round(2),
        f"all_platform_equal_{len(all_platforms)}平台等权%": all_equal.round(2),
    })
    comp.index.name = "情感"
    return comp, elig, eligible


# ============================================================
# 五、模型与人工一致率
# ============================================================
def consistency_stats(df):
    rev = df[df["manual_reviewed"] == "是"].copy()
    agree_all = (rev["manual_sentiment"] == rev["sentiment_pred"]).sum()
    rate_all = agree_all / len(rev) * 100
    rev_nr = rev[rev["manual_sentiment"] != "other"]
    agree_nr = (rev_nr["manual_sentiment"] == rev_nr["sentiment_pred"]).sum()
    rate_nr = agree_nr / len(rev_nr) * 100 if len(rev_nr) else np.nan

    summary = pd.DataFrame({
        "口径": ["全部复核记录(含 other)", "排除 other 后"],
        "记录数": [len(rev), len(rev_nr)],
        "一致数": [agree_all, agree_nr],
        "一致率%": np.round([rate_all, rate_nr], 2),
    })

    rev["_agree"] = rev["manual_sentiment"] == rev["sentiment_pred"]
    by_pred = rev.groupby("sentiment_pred")["_agree"].mean().reset_index()
    by_pred.columns = ["模型标签", "人工确认率%"]
    by_pred["人工确认率%"] = (by_pred["人工确认率%"] * 100).round(2)
    rev = rev.drop(columns="_agree")

    cross = pd.crosstab(rev["sentiment_pred"], rev["manual_sentiment"]).reindex(
        index=["positive", "neutral", "negative"], columns=EMOTIONS + ["other"], fill_value=0)
    cross.insert(0, "模型标签", cross.index)
    return summary, by_pred, cross


# ============================================================
# 六、写汇总 xlsx
# ============================================================
def _preview(s):
    s = s.astype(str)
    return s.map(lambda t: t[:CONTENT_PREVIEW] + ("…" if len(t) > CONTENT_PREVIEW else ""))


def write_excel(df, comp, elig, rev_stats, by_pred, cross, label):
    from openpyxl.utils import get_column_letter

    def _sheet(writer, name, data, note=None):
        data.to_excel(writer, sheet_name=name, index=False, startrow=0)
        ws = writer.sheets[name]
        if note:
            ws.insert_rows(1)
            ws.cell(row=1, column=1, value=note)
        for col, c in enumerate(data.columns, 1):
            w = max(data[c].map(lambda x: len(str(x))).max(), len(str(c))) if len(data) else len(str(c))
            ws.column_dimensions[get_column_letter(col)].width = min(max(w + 2, 10), 60)

    main = df[df["sentiment_final"] != "other"]
    n_all = len(df)
    n_main = len(main)
    n_other = n_all - n_main
    n_review = int((df["manual_reviewed"] == "是").sum())
    n_eligible = int(elig["纳入均衡"].sum())
    eligible_list = elig.loc[elig["纳入均衡"], "平台"].tolist()

    detail_cols = ["record_id", "anime_name", "form", "source", "lang_group",
                   "content_length", "sentiment_pred", "sentiment_final"]

    with pd.ExcelWriter(OUT_XLSX, engine="openpyxl") as writer:
        # ---- 表1 总览 ----
        overview = pd.DataFrame({
            "指标": [
                f"层级", "全量记录数", "排除 other 后主样本数", "other 记录数",
                "人工复核数", "复核覆盖率%",
                "— pooled 口径 —",
                "— platform_balanced（正文主口径）—",
                "— all_platform_equal（敏感性）—",
                "准入平台",
                "未准入平台",
            ],
            "数值": [
                label, n_all, n_main, n_other,
                n_review, np.round(n_review / n_all * 100, 2),
                " / ".join(f"{e}:{v}%" for e, v in zip(comp.index, comp.iloc[:, 0])),
                " / ".join(f"{e}:{v}%" for e, v in zip(comp.index, comp.iloc[:, 1])),
                " / ".join(f"{e}:{v}%" for e, v in zip(comp.index, comp.iloc[:, 2])),
                ", ".join(eligible_list) if eligible_list else "(无)",
                ", ".join(elig.loc[~elig["纳入均衡"], "平台"].tolist()) or "(无)",
            ],
        })
        _sheet(writer, "总览", overview, note=f"{label} 层三口径汇总；主比例均排除 other。")

        # ---- 表2 三口径对比 ----
        comp_out = comp.reset_index()
        _sheet(writer, "三口径对比", comp_out,
               note="platform_balanced 是正文主口径；all_platform_equal 仅敏感性分析。")

        # ---- 表3 平台准入 ----
        _sheet(writer, "平台准入", elig,
               note=f"准入规则：样本量 n>={MIN_PLATFORM_N} 且 覆盖作品数>={MIN_PLATFORM_WORKS}（在主样本上计算）。")

        # ---- 表4 最终情感分布（pooled 计数） ----
        vc_all = df["sentiment_final"].value_counts()
        vc_main = main["sentiment_final"].value_counts()
        rows = []
        for e in vc_all.index:
            main_pct = np.round(vc_main.get(e, 0) / n_main * 100, 2) if e != "other" else None
            rows.append({"情感": e, "计数": int(vc_all[e]),
                         "全量占比%": np.round(vc_all[e] / n_all * 100, 2),
                         "主样本占比%(排除other)": main_pct})
        _sheet(writer, "最终情感分布", pd.DataFrame(rows))

        # ---- 表5 平台×情感 ----
        pv = main.pivot_table(index="source", columns="sentiment_final",
                              values="record_id", aggfunc="count", fill_value=0)
        pv = pv.reindex(columns=EMOTIONS, fill_value=0)
        pr = pv.div(pv.sum(axis=1), axis=0).fillna(0) * 100
        pm = pv.copy()
        for e in EMOTIONS:
            pm[f"{e}_占比%"] = pr[e].round(2)
        pm.insert(0, "平台", pm.index)
        pm.insert(1, "总计数", pm[EMOTIONS].sum(axis=1).astype(int))
        pm = pm[["平台", "总计数"] + EMOTIONS + [f"{e}_占比%" for e in EMOTIONS]]
        _sheet(writer, "平台×情感", pm, note="仅主样本；占比为平台内行占比。")

        # ---- 表6 作品×情感 ----
        av = main.pivot_table(index="anime_name", columns="sentiment_final",
                              values="record_id", aggfunc="count", fill_value=0)
        av = av.reindex(columns=EMOTIONS, fill_value=0)
        ar = av.div(av.sum(axis=1), axis=0).fillna(0) * 100
        am = av.copy()
        for e in EMOTIONS:
            am[f"{e}_占比%"] = ar[e].round(2)
        am.insert(0, "作品", am.index)
        am.insert(1, "总计数", am[EMOTIONS].sum(axis=1).astype(int))
        am = am[["作品", "总计数"] + EMOTIONS + [f"{e}_占比%" for e in EMOTIONS]]
        _sheet(writer, "作品×情感", am, note="仅主样本（排除 other）。")

        # ---- 表7 平台×作品×情感 ----
        spv = main.pivot_table(index=["source", "anime_name"], columns="sentiment_final",
                               values="record_id", aggfunc="count", fill_value=0)
        spv = spv.reindex(columns=EMOTIONS, fill_value=0)
        spv.insert(0, "平台×作品", spv.index.map(lambda t: f"{t[0]} | {t[1]}"))
        spv.insert(1, "总计数", spv[EMOTIONS].sum(axis=1).astype(int))
        spv = spv[["平台×作品", "总计数"] + EMOTIONS]
        _sheet(writer, "平台×作品×情感", spv, note="仅主样本（排除 other）。")

        # ---- 表8 语言×情感 ----
        lv = main.pivot_table(index="lang_group", columns="sentiment_final",
                              values="record_id", aggfunc="count", fill_value=0)
        lv = lv.reindex(columns=EMOTIONS, fill_value=0)
        lr = lv.div(lv.sum(axis=1), axis=0).fillna(0) * 100
        lm = lv.copy()
        for e in EMOTIONS:
            lm[f"{e}_占比%"] = lr[e].round(2)
        lm.insert(0, "语言", lm.index)
        lm.insert(1, "总计数", lm[EMOTIONS].sum(axis=1).astype(int))
        lm = lm[["语言", "总计数"] + EMOTIONS + [f"{e}_占比%" for e in EMOTIONS]]
        _sheet(writer, "语言×情感", lm, note="仅主样本（排除 other）。")

        # ---- 表9 人工复核统计 ----
        grp = df[df["manual_reviewed"] == "是"]
        review_stats = pd.DataFrame({
            "项目": [
                "人工复核总条数", "核心复核清单", "补充抽审清单",
                "manual=positive", "manual=neutral", "manual=negative",
                "manual=mixed", "manual=other", "manual_note 非空",
            ],
            "数值": [
                len(grp),
                int((grp["review_queue_group"] == "核心复核清单").sum()),
                int((grp["review_queue_group"] == "补充抽审清单").sum()),
                int((grp["manual_sentiment"] == "positive").sum()),
                int((grp["manual_sentiment"] == "neutral").sum()),
                int((grp["manual_sentiment"] == "negative").sum()),
                int((grp["manual_sentiment"] == "mixed").sum()),
                int((grp["manual_sentiment"] == "other").sum()),
                int((grp["manual_note"].astype(str).str.strip() != "").sum()),
            ],
        })
        _sheet(writer, "人工复核统计", review_stats)

        # ---- 表10 一致率 ----
        _sheet(writer, "模型与人工一致率", rev_stats, note="一致率只在人工复核的记录内计算。")
        _sheet(writer, "分标签一致率", by_pred)
        cross_out = cross.copy()
        cross_out.insert(1, "合计", cross_out[EMOTIONS + ["other"]].sum(axis=1).astype(int))
        _sheet(writer, "模型×人工交叉表", cross_out, note="行为模型 sentiment_pred，列为人工 manual_sentiment。")

        # ---- 表11 标签修改明细 ----
        changed = df[(df["manual_reviewed"] == "是") & (df["label_changed"] == "是")].copy()
        if len(changed):
            changed = changed.sort_values(["manual_sentiment", "record_id"])
            changed_out = changed[detail_cols + ["manual_sentiment", "manual_note"]].copy()
            changed_out["content_clean(预览)"] = _preview(changed["content_clean"])
        else:
            changed_out = pd.DataFrame({"提示": ["没有人工与模型不一致的记录"]})
        _sheet(writer, "标签修改明细", changed_out, note=f"共 {len(changed)} 条人工与模型不一致。")

        # ---- 表12 other 记录明细 ----
        other_rows = df[df["sentiment_final"] == "other"].copy()
        if len(other_rows):
            other_out = other_rows[detail_cols].copy()
            other_out["manual_note"] = other_rows["manual_note"]
            other_out["content_clean(预览)"] = _preview(other_rows["content_clean"])
        else:
            other_out = pd.DataFrame({"提示": ["没有 other 记录"]})
        _sheet(writer, "other记录明细", other_out,
               note=f"共 {len(other_rows)} 条 other，保留在数据中但不纳入主比例分母。")

    print(f"[6/7] 汇总已写入 {OUT_XLSX}")


# ============================================================
# 七、主流程
# ============================================================
def main():
    global OUT_XLSX
    parser = argparse.ArgumentParser(description="情感复核回填 + 三口径汇总")
    parser.add_argument("--csv", default="03_sentiment_all.csv")
    parser.add_argument("--xlsx", default="sentiment_review_core.xlsx")
    parser.add_argument("--out-csv", default="04_sentiment_final.csv")
    parser.add_argument("--out-xlsx", default="sentiment_final_summary.xlsx")
    parser.add_argument("--label", default="")
    parser.add_argument("--lenient", action="store_true")
    args = parser.parse_args()
    OUT_XLSX = args.out_xlsx

    print(f"[{args.label or '数据'}] 读取全量基线 {args.csv}")
    df = pd.read_csv(args.csv, encoding="utf-8-sig", dtype=str, keep_default_na=False)
    print(f"      共 {len(df)} 条, {df.shape[1]} 列")

    review = load_review_tables(args.xlsx)
    validate_review(review, df, lenient=args.lenient)

    print("[3/7] 回填人工结果并生成 sentiment_final")
    out = merge_back(df, review)
    out = finalize(out)

    print("[4/7] 计算三口径（pooled / platform_balanced / all_platform_equal）")
    main_df = out[out["sentiment_final"] != "other"]
    comp, elig, eligible = compute_three_views(main_df)
    print("      三口径对比:")
    print(comp.to_string())
    print("      平台准入:")
    print(elig.to_string(index=False))

    print("[5/7] 计算模型与人工一致率")
    rev_stats, by_pred, cross = consistency_stats(out)
    print(rev_stats.to_string(index=False))

    out.to_csv(args.out_csv, index=False, encoding="utf-8-sig")
    print(f"[6/7] 最终数据已写入 {args.out_csv}")

    write_excel(out, comp, elig, rev_stats, by_pred, cross, args.label)
    print(f"[7/7] 完成。人工复核 {int((out['manual_reviewed']=='是').sum())} 条, "
          f"other {int((out['sentiment_final']=='other').sum())} 条, "
          f"准入平台 {eligible}。")


if __name__ == "__main__":
    main()
